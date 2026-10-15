"""Build editable, illustrated Word manuscripts from the paper Markdown sources.

Optional document dependencies: python-docx and CairoSVG (not app dependencies).
"""

from __future__ import annotations

import argparse
import io
import re
from pathlib import Path
from typing import Any

import cairosvg
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from docx.shared import Cm, Pt, RGBColor

#: Plain mode: Times New Roman 12 throughout, black, no header, footer, rules,
#: shading or accent colour. Journals generally want a manuscript that carries no
#: styling of its own, so their own typesetting is what the reader sees.
PLAIN = False
PLAIN_FONT, PLAIN_SIZE = "Times New Roman", Pt(12)


def inline(paragraph: Any, text: str) -> None:
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|https?://\S+|\[\d+(?:[,– -]+\d+)*\])", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            if PLAIN:
                run.font.name, run.font.size = PLAIN_FONT, PLAIN_SIZE
            else:
                run.font.name, run.font.size = "Consolas", Pt(9)
        elif token.startswith(("http://", "https://")):
            url = token.rstrip(".,;")
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "000000" if PLAIN else "166C62")
            properties.append(color)
            run.append(properties)
            value = OxmlElement("w:t")
            value.text = url
            run.append(value)
            link.append(run)
            paragraph._p.append(link)
            paragraph.add_run(token[len(url) :])
        elif re.fullmatch(r"\[\d+(?:[,– -]+\d+)*\]", token):
            for part in re.split(r"(\d+)", token):
                if part.isdigit():
                    link = OxmlElement("w:hyperlink")
                    link.set(qn("w:anchor"), f"ref{part}")
                    run = OxmlElement("w:r")
                    value = OxmlElement("w:t")
                    value.text = part
                    run.append(value)
                    link.append(run)
                    paragraph._p.append(link)
                else:
                    paragraph.add_run(part)
        else:
            paragraph.add_run(token)


def add_table(document: Any, lines: list[str]) -> None:
    rows = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    weights = [max(1.5, min(3.0, len(text) ** 0.5 / 2)) for text in rows[0]]
    widths = [weight * 16.6 / sum(weights) for weight in weights]
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Cm(width)
    for number, (row, contents) in enumerate(zip(table.rows, rows, strict=True)):
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if number == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for index, (cell, content) in enumerate(zip(row.cells, contents, strict=True)):
            cell.width = Cm(widths[index])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0 if PLAIN else 1.05
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = paragraph.add_run(content)
            run.font.size = PLAIN_SIZE if PLAIN else Pt(8.5 if count >= 7 else 9)
            if number == 0:
                run.bold = True
                if not PLAIN:
                    run.font.color.rgb = RGBColor.from_string("FFFFFF")
            shading = OxmlElement("w:shd")
            shading.set(
                qn("w:fill"),
                "FFFFFF"
                if PLAIN
                else ("166C62" if number == 0 else ("EDF5F3" if number % 2 else "FFFFFF")),
            )
            cell._tc.get_or_add_tcPr().append(shading)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(document: Any, source: Path, number: str, caption: str) -> None:
    if number not in {"1", "S1"}:
        document.add_page_break()
    raster = cairosvg.svg2png(url=str(source), output_width=2600)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = paragraph.add_run().add_picture(io.BytesIO(raster), width=Cm(16.6))
    picture._inline.docPr.set("descr", caption)
    nsmap["asvg"] = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
    vector = Part(
        PackURI(f"/word/media/figure{number}.svg"),
        "image/svg+xml",
        source.read_bytes(),
        document.part.package,
    )
    relationship = document.part.relate_to(vector, RT.IMAGE)
    blip = picture._inline.xpath(".//a:blip")[0]
    extensions = OxmlElement("a:extLst")
    extension = OxmlElement("a:ext")
    extension.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    svg = OxmlElement("asvg:svgBlip")
    svg.set(qn("r:embed"), relationship)
    extension.append(svg)
    extensions.append(extension)
    blip.append(extensions)
    inline(document.add_paragraph(style="Caption"), caption)


def build(source: Path, out: Path) -> Path:
    if out.exists():
        raise ValueError("Refusing to overwrite an existing Word manuscript")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.2)
    section.header_distance = section.footer_distance = Cm(1)
    normal = document.styles["Normal"]
    if PLAIN:
        normal.font.name, normal.font.size = PLAIN_FONT, PLAIN_SIZE
        normal.font.color.rgb = RGBColor.from_string("000000")
        # Double spacing, which most journals ask for so reviewers can annotate.
        normal.paragraph_format.line_spacing = 2.0
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0)
    else:
        normal.font.name, normal.font.size = "Calibri", Pt(11)
        normal.paragraph_format.line_spacing = 1.10
        normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    for name, size in [("Title", 20), ("Heading 1", 15), ("Heading 2", 12)]:
        style = document.styles[name]
        if PLAIN:
            # Headings differ from body text by weight alone, not size or colour.
            style.font.name, style.font.size = PLAIN_FONT, PLAIN_SIZE
            style.font.color.rgb = RGBColor.from_string("000000")
            style.font.bold = True
            style.paragraph_format.line_spacing = 2.0
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(0)
        else:
            style.font.name, style.font.size = "Calibri", Pt(size)
            style.font.color.rgb = RGBColor.from_string("166C62")
            style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
            style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    caption = document.styles["Caption"]
    if PLAIN:
        caption.font.name, caption.font.size = PLAIN_FONT, PLAIN_SIZE
        caption.font.color.rgb = RGBColor.from_string("000000")
        caption.font.italic = False
        caption.paragraph_format.line_spacing = 2.0
    else:
        caption.font.size = Pt(9)
        caption.font.color.rgb = RGBColor.from_string("425564")

    if not PLAIN:
        header = section.header.paragraphs[0]
        header.text = "SARscope · Manuscript draft"
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string("6A7784")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    lines = source.read_text().splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("#"):
            marks, title = line.split(" ", 1)
            paragraph = document.add_paragraph(
                style="Title" if marks == "#" else f"Heading {len(marks) - 1}"
            )
            inline(paragraph, title)
            if title in {"Figure legends", "References"}:
                paragraph.paragraph_format.page_break_before = True
            index += 1
            continue
        if line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            add_table(document, block)
            continue
        block = [line]
        index += 1
        while (
            index < len(lines)
            and lines[index].strip()
            and not lines[index].startswith(("#", "|", "- "))
            and not re.match(r"^\d+\.\s", lines[index])
        ):
            block.append(lines[index].strip())
            index += 1
        text = " ".join(block)
        match = re.match(r"^Figure (S?\d+): `([^`]+)`\. (.*)", text)
        if match:
            number, relative, prose = match.groups()
            add_figure(document, source.parent / relative, number, f"Figure {number}. {prose}")
        else:
            paragraph = document.add_paragraph(
                style="List Bullet" if text.startswith("- ") else None
            )
            inline(paragraph, text[2:] if text.startswith("- ") else text)
            if re.match(r"^\d+\. ", text):
                paragraph.paragraph_format.left_indent = Cm(0.5)
                paragraph.paragraph_format.first_line_indent = Cm(-0.5)
                number = text.split(".", 1)[0]
                start = OxmlElement("w:bookmarkStart")
                start.set(qn("w:id"), number)
                start.set(qn("w:name"), f"ref{number}")
                end = OxmlElement("w:bookmarkEnd")
                end.set(qn("w:id"), number)
                paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
                paragraph._p.append(end)
    document.core_properties.title = lines[0].lstrip("# ")
    document.core_properties.subject = (
        "Retrospective positional SAR study; draft pending author review"
    )
    document.core_properties.author = "Yassir Boulaamane"
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    return out


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument(
        "--plain",
        action="store_true",
        help="Times New Roman 12, all black, no header, footer, shading or accent colour",
    )
    args = parser.parse_args()
    globals()["PLAIN"] = args.plain
    print(build(args.source, args.out))
